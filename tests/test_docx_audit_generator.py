# -*- coding: utf-8 -*-
"""
Test Suite Kiểm Thử Nghiệm Thu Động Cơ Chế Bản Hồ Sơ Kiểm Toán (.docx)
Định vị: tests/test_docx_audit_generator.py
Áp dụng: ENGINEERING_DISCIPLINE_WORKFLOW.md [PHA 3: BƯỚC 3.3 & BƯỚC 3.5]
"""

import os
import re
import sys
import tempfile
import pytest

# Đảm bảo đường dẫn gốc của dự án nằm trong sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import docx

from scripts.docx_audit_generator import (
    ImageAssetLink,
    SectionAuditItem,
    ChapterAuditPayload,
    DocxAuditResult,
    DriveFSSqliteExtractor,
    DocxAuditGenerator,
    generate_chapter_audit_docx,
    auto_generate_audit_docx_for_chapter,
)


# ===========================================================================
# Regex từ spark-visual-cockpit-v2.0.0/content.js để đối soát độc lập
# ===========================================================================

# Regex trích xuất Drive File ID (dòng 867-871)
RE_DRIVE_FILE_ID = re.compile(r"/(?:file/)?d/([a-zA-Z0-9_-]{15,})|[?&]id=([a-zA-Z0-9_-]{15,})", re.I)

# Regex phát hiện ảnh/đồ thị cho Tab Trái Canvas (dòng 1094 & 1209)
RE_DIAGRAM_FILE = re.compile(r"(?:^|\n|\-\s*|\*\s*)(\[(?:Goc|Mo_Phong|Original|Simulation)\][^\s:]+)\s*:", re.I)

# Regex phát hiện PDF cho Tab Phải Reader (dòng 1133, 1239, 1806, 1834)
RE_PDF_DOC = re.compile(
    r"(?:bài giảng|giáo trình|tài liệu|lecture|slides?|sách|văn bản|tập tài liệu|chương \d|phân mục \d|bản chuyên khảo|chuyên khảo|mục \d)(?:\.pdf)?\b",
    re.I
)

# Regex trích xuất tiêu đề PDF (dòng 1145, 1244, 1811, 1836)
RE_PDF_TITLE = re.compile(
    r"(?:Tài liệu|Bài giảng|Giáo trình|Bản Chuyên khảo|Chuyên khảo|Phân mục|Mục|Chương)\s*[:\-\s]*([^:\n|]+)",
    re.I
)


def extract_drive_id(url: str) -> str:
    m = RE_DRIVE_FILE_ID.search(url)
    if m:
        return m.group(1) or m.group(2)
    return ""


# ===========================================================================
# Unit Tests & Acceptance Tests
# ===========================================================================

def test_drivefs_sqlite_extractor_detection():
    """Kiểm tra việc phát hiện và kết nối Read-Only tới các DB DriveFS trên hệ thống."""
    db_paths = DriveFSSqliteExtractor.get_drivefs_db_paths()
    assert isinstance(db_paths, list), "Đầu ra danh sách đường dẫn DB phải là một List"
    if db_paths:
        for p in db_paths:
            assert os.path.exists(p), f"Tệp DB không tồn tại: {p}"
            # Thử truy vấn một tệp không tồn tại để kiểm tra không bị khóa hay văng crash
            res = DriveFSSqliteExtractor.query_file_id("non_existent_file_xyz_12345.pdf", max_retries=1)
            assert res is None, "Tệp không tồn tại phải trả về None"


def test_docx_creation_valid_openxml():
    """Kiểm tra việc tạo tệp .docx chuẩn OpenXML và tính toàn vẹn nhị phân."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        target_path = os.path.join(tmp_dir, "test_audit.docx")

        sample_sections = [
            SectionAuditItem(
                section_id="1.1",
                section_title="Cơ sở Bán dẫn và Tiếp giáp p-n",
                pdf_filename="Chuyen_Khao_Laser_Diode_Muc_1.1.pdf",
                pdf_drive_id="1pVUOBjeA_MVRpYQ2N2Gviq7NKU2ZrdD2",
                mpt_count=3,
                images=[
                    ImageAssetLink(
                        filename="[Mo_Phong]_Muc_1.1_Dai_Nang_Luong.png",
                        image_type="Mo_Phong",
                        drive_file_id="1jI6QSIn2rwAbKZoxMlkI6Hz_DprEBJrJ"
                    ),
                    ImageAssetLink(
                        filename="[Goc]_Figure_1.1_Cau_Truc_Tiep_Giap.png",
                        image_type="Goc",
                        drive_file_id="1bxb6hVTb6SDBeTX6mzUaeoixaRxBKwns"
                    )
                ]
            )
        ]

        payload = ChapterAuditPayload(
            subject_name="Laser Diode - Cấu trúc và Nguyên lý",
            chapter_id="1",
            chapter_title="Chương 1: Cơ sở Bán dẫn",
            sections=sample_sections,
            target_drive_root=tmp_dir
        )

        result = generate_chapter_audit_docx(payload, custom_target_path=target_path)

        assert os.path.exists(result.docx_target_path), "File docx mục tiêu không tồn tại sau khi tạo"
        assert result.file_size_bytes > 0, "File docx rỗng 0 bytes"
        assert result.file_size_bytes < 100 * 1024, "Dung lượng docx vượt quá 100 KB (vi phạm yêu cầu siêu nhẹ)"
        assert result.total_sections == 1
        assert result.total_images == 2
        assert result.resolved_pdf_links == 1
        assert result.resolved_image_links == 2
        assert result.status == "SUCCESS"

        # Kiểm tra mở lại bằng thư viện docx (chứng minh XML không hỏng)
        doc = docx.Document(result.docx_target_path)
        assert len(doc.tables) >= 1, "Tài liệu phải chứa ít nhất 1 bảng biểu"
        assert len(doc.paragraphs) > 5, "Tài liệu phải có đủ tiêu đề và nội dung chi tiết"


def test_content_js_regex_compatibility_100_percent():
    """
    Đọc ngược toàn bộ văn bản từ file .docx và chạy qua chính các regex của
    spark-visual-cockpit-v2.0.0/content.js để khẳng định tỷ lệ khớp 100%.
    """
    with tempfile.TemporaryDirectory() as tmp_dir:
        target_path = os.path.join(tmp_dir, "test_regex_audit.docx")

        sections = [
            SectionAuditItem(
                section_id="2.1",
                section_title="Cấu trúc Diode Dị thể Kép DH",
                pdf_filename="Chuyen_Khao_Laser_Diode_Muc_2.1.pdf",
                pdf_drive_id="1W_F2InCiKejObVA5Nv6croogPZ840oM5",
                mpt_count=4,
                images=[
                    ImageAssetLink(
                        filename="[Mo_Phong]_Muc_2.1_Phan_Bo_Mode_Quang.png",
                        image_type="Mo_Phong",
                        drive_file_id="1xGVKtPAzQ7Zp5M6PS4Oy75PScCR20Wm6"
                    ),
                    ImageAssetLink(
                        filename="[Goc]_Figure_2.1_Gian_do_dai_nang_luong.png",
                        image_type="Goc",
                        drive_file_id="1s98ftDa1_KokziU8alDTjoqEZgEEO2Ll"
                    )
                ]
            )
        ]

        payload = ChapterAuditPayload(
            subject_name="Laser Diode",
            chapter_id="2",
            chapter_title="Chương 2",
            sections=sections,
            target_drive_root=tmp_dir
        )

        res = generate_chapter_audit_docx(payload, custom_target_path=target_path)
        doc = docx.Document(res.docx_target_path)

        # Trích xuất tất cả các đoạn văn bản
        raw_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # 1. Kiểm tra nhận diện Link PDF
        pdf_lines = [l for l in raw_lines if "Bản Chuyên khảo" in l]
        assert len(pdf_lines) == 1, "Phải có đúng 1 dòng Bản Chuyên khảo"
        pdf_line = pdf_lines[0]

        # Kiểm tra regex RE_PDF_DOC từ content.js
        assert RE_PDF_DOC.search(pdf_line), f"Dòng PDF không khớp RE_PDF_DOC: '{pdf_line}'"

        # Kiểm tra trích xuất tiêu đề PDF
        title_m = RE_PDF_TITLE.search(pdf_line)
        assert title_m, f"Không trích xuất được tiêu đề PDF từ: '{pdf_line}'"
        assert "Mục 2.1" in title_m.group(0)

        # Kiểm tra Drive URL trong dòng PDF
        pdf_drive_id = extract_drive_id(pdf_line)
        assert pdf_drive_id == "1W_F2InCiKejObVA5Nv6croogPZ840oM5", f"Drive ID trích xuất sai: {pdf_drive_id}"

        # 2. Kiểm tra nhận diện Hình ảnh (Canvas Tab Trái)
        img_lines = [l for l in raw_lines if l.startswith("- [")]
        assert len(img_lines) == 2, f"Phải có 2 dòng ảnh, thực tế có: {len(img_lines)}"

        for line in img_lines:
            diag_m = RE_DIAGRAM_FILE.search(line)
            assert diag_m, f"Dòng ảnh không khớp RE_DIAGRAM_FILE: '{line}'"
            matched_filename = diag_m.group(1)
            assert matched_filename in [
                "[Mo_Phong]_Muc_2.1_Phan_Bo_Mode_Quang.png",
                "[Goc]_Figure_2.1_Gian_do_dai_nang_luong.png"
            ]

            img_drive_id = extract_drive_id(line)
            assert len(img_drive_id) >= 15, f"Drive ID hình ảnh không hợp lệ: '{line}'"


def test_pending_sync_graceful_fallback():
    """Kiểm tra cơ chế xử lý tệp chưa kịp đồng bộ ID (fallback PENDING_SYNC)."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        target_path = os.path.join(tmp_dir, "test_pending.docx")

        # Mục chưa có drive_file_id và tên file không có trong SQLite
        sections = [
            SectionAuditItem(
                section_id="9.9",
                section_title="Mục Giả Lập Chưa Đồng Bộ",
                pdf_filename="Chuyen_Khao_Chua_Tung_Ton_Tai_9.9.pdf",
                images=[
                    ImageAssetLink(
                        filename="[Mo_Phong]_Chua_Sync_9.9.png",
                        image_type="Mo_Phong"
                    )
                ]
            )
        ]

        payload = ChapterAuditPayload(
            subject_name="Test Subject",
            chapter_id="9",
            chapter_title="Chương 9",
            sections=sections,
            target_drive_root=tmp_dir
        )

        res = generate_chapter_audit_docx(payload, custom_target_path=target_path)
        assert res.status == "PARTIAL_SYNC"
        assert res.pending_sync_count == 2  # Cả PDF và ảnh đều pending
        assert os.path.exists(res.docx_target_path)

        doc = docx.Document(res.docx_target_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "PENDING_SYNC" in full_text
        assert "Mục 9.9" in full_text


def test_zero_thumbnail_binary_invariant():
    """Kiểm tra ràng buộc thép: Tuyệt đối không nhúng ảnh thumbnail nhị phân vào docx."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        target_path = os.path.join(tmp_dir, "test_no_binary.docx")

        sections = [
            SectionAuditItem(
                section_id="1.1",
                section_title="Mục Kiểm Tra Binary",
                pdf_filename="Chuyen_Khao_Laser_Diode_Muc_1.1.pdf",
                pdf_drive_id="1pVUOBjeA_MVRpYQ2N2Gviq7NKU2ZrdD2",
                images=[
                    ImageAssetLink(
                        filename="[Mo_Phong]_Test_Image.png",
                        image_type="Mo_Phong",
                        drive_file_id="1jI6QSIn2rwAbKZoxMlkI6Hz_DprEBJrJ"
                    )
                ]
            )
        ]

        payload = ChapterAuditPayload(
            subject_name="Laser Diode",
            chapter_id="1",
            chapter_title="Chương 1",
            sections=sections,
            target_drive_root=tmp_dir
        )

        res = generate_chapter_audit_docx(payload, custom_target_path=target_path)
        doc = docx.Document(res.docx_target_path)

        # Kiểm tra danh sách image parts trong gói OpenXML
        image_parts = [p for p in doc.part.related_parts.values() if "image" in p.content_type]
        assert len(image_parts) == 0, f"Vi phạm: Có {len(image_parts)} ảnh nhị phân bị nhúng vào docx!"


def test_live_chapter_2_drive_audit_generation():
    r"""Nghiệm thu thực tế trên dữ liệu Google Google Drive của Chương 2 Laser Diode."""
    res = auto_generate_audit_docx_for_chapter("Laser Diode - Cấu trúc và Nguyên lý", "2")
    assert res.status == "SUCCESS"
    assert res.total_sections == 8
    assert res.total_images == 26
    assert res.resolved_pdf_links == 8
    assert res.resolved_image_links == 26
    assert res.pending_sync_count == 0
    assert os.path.exists(res.docx_target_path)
    assert res.file_size_bytes > 30 * 1024
    assert res.file_size_bytes < 100 * 1024
