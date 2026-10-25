# -*- coding: utf-8 -*-
"""
Test Suite Kiểm Chứng Khóa Chốt Chặn Đối Chiếu Chéo Ground Truth (Cross-Audit Interlock)
Định vị: tests/test_ground_truth_interlock.py
Áp dụng: ENGINEERING_DISCIPLINE_WORKFLOW.md [PHA 1 & PHA 3: BƯỚC 3.3]
"""

import os
import sys
import json
import hashlib
import tempfile
import pytest

# Thiết lập đường dẫn import
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOT_AGENTS = os.path.join(PROJECT_ROOT, ".agents")
for p in [DOT_AGENTS, PROJECT_ROOT]:
    if p not in sys.path:
        sys.path.insert(0, p)

from schemas.blueprint import (
    CurriculumBlueprint,
    KnowledgeDAGNode,
    MicroPointTopicSpec,
    TerminologyEntry,
)
from schemas.manifest import (
    VisualAssetManifest,
    VisualAssetItem,
    TripleBufferMetrics,
    OCRChainPedagogicalCaption,
)
from schemas.payload import (
    DraftPayload,
    MicroPointTopicPayload,
    TeleologySpatialContext,
    MechanisticBlendingEquation,
    InternalASTLinterReport,
    ToulminFourBeatPipeline,
    OrthogonalAndCERBlock,
    OrthogonalPerspective,
    OpenCERChallenge,
)
from schemas.audit import (
    StationResult,
    DOHGStationsResult,
    PublicationMetadata,
)
from agents.quality_gatekeeper.gatekeeper import QualityGatekeeper


def _create_dummy_image_with_sidecar(dir_path: str, filename: str, content: bytes) -> str:
    """Tạo file ảnh giả lập và sidecar manifest tương ứng."""
    img_path = os.path.join(dir_path, filename)
    with open(img_path, "wb") as f:
        f.write(content)

    sidecar_path = os.path.splitext(img_path)[0] + ".manifest.json"
    sidecar_data = {
        "sha256": hashlib.sha256(content).hexdigest(),
        "geometry_audit": {
            "status": "PASS",
            "label_overlaps": 0,
        },
        "zero_slash_fractions": "PASS",
        "canvas_width_inch": 6.5,
    }
    with open(sidecar_path, "w", encoding="utf-8") as f:
        json.dump(sidecar_data, f, ensure_ascii=False, indent=2)

    return img_path


def _create_minimal_payload(section_id: str = "MUC_3_1") -> DraftPayload:
    """Tạo DraftPayload hợp lệ tối thiểu."""
    ortho_cer = OrthogonalAndCERBlock(
        orthogonal_perspective=OrthogonalPerspective(
            dimension_shifts=[{"axis": "z vs k", "insight": "Đối trọng dải năng lượng"}],
            intuitive_dissection="Giải phẫu góc nhìn đối trọng",
            practical_limitations="Hạn chế thực tế màng mỏng",
        ),
        open_cer_challenge=OpenCERChallenge(
            challenge_context="Bối cảnh phản ứng quang xúc tác",
            diagnostic_question="Cơ chế tách điện tích tại dị liên kết Z-scheme?",
            claim_guidance="Khẳng định hướng truyền electron",
            evidence_required="Thế dải dẫn ECB và thế dải hóa trị EVB",
            reasoning_rubric="Lập luận theo cơ chế tái hợp trung gian",
            misconception_trap="Nhầm lẫn giữa Type-II dị thể và Direct Z-scheme",
            chat_evaluation_key=None,
        ),
    )

    mpt_payload = MicroPointTopicPayload(
        mpt_id="MPT_3_1_1",
        mpt_title="Cơ chế quang xúc tác phân rã chất màu",
        visual_anchor_id="FIG_3_1_1",
        teleology_and_spatial_context=TeleologySpatialContext(
            systemic_teleology_sbf="Giải quyết bài toán xử lý nước thải",
            plain_english_axiom="Quang sinh cặp electron-lỗ trống",
            spatial_topology_anchor="Dải năng lượng dị liên kết dị thể",
            cognitive_handover_to_block_2="Hình thái dải năng lượng và động học truyền điện tích",
        ),
        mechanistic_blending_equation=MechanisticBlendingEquation(
            first_principles_origin="Phương trình bảo toàn điện tích và động học phản ứng",
            pillar_equation_latex="r = -\\frac{dC}{dt} = k_{app} C",
            visual_evidence="Dữ liệu phổ hấp thụ và giản đồ XRD",
            micro_dynamics_warrant="Lỗ trống quang sinh oxy hóa trực tiếp chất màu",
            math_blending_scaling="Động học biểu kiến bậc 1 theo mô hình Langmuir-Hinshelwood",
            engineering_claim="Tăng tốc độ phản ứng quang xúc tác gấp 3 lần",
            four_beat_pipeline=ToulminFourBeatPipeline(
                beat_1_visual_evidence="Dữ liệu phổ hấp thụ và giản đồ XRD",
                beat_2_micro_dynamics_warrant="Lỗ trống quang sinh oxy hóa trực tiếp chất màu",
                beat_3_math_blending_scaling="Động học biểu kiến bậc 1 theo mô hình Langmuir-Hinshelwood",
                beat_4_engineering_claim="Tăng tốc độ phản ứng quang xúc tác gấp 3 lần",
            ),
        ),
        block_3_orthogonal_cer=ortho_cer,
    )

    return DraftPayload(
        section_id=section_id,
        section_title="Quang xúc tác xử lý môi trường",
        cod_rounds_completed=3,
        mpt_payloads=[mpt_payload],
        linter_pre_check=InternalASTLinterReport(
            passed=True,
            raw_latex_leaks_count=0,
            slash_fraction_count=0,
            raw_subscript_count=0,
            unwarranted_claims_count=0,
        ),
    )


def test_01_reproduce_omission_failure():
    """
    [BÀI TEST TÁI HIỆN LỖI - BẮT BUỘC FAIL BAN ĐẦU]:
    Mô phỏng trường hợp Mục 3.1 Quang xúc tác:
    - Ground Truth chứa 3 thực thể: GT_SCHEME_1, GT_FIG_1_A, GT_FIG_1_B.
    - Manifest chỉ bóc tách 1 ảnh gộp: FIG_1 (thiếu Scheme 1 và Fig 1b).
    - Blueprint chỉ có MPT cho Fig 1 (thiếu MPT cho Scheme 1).
    KẾT QUẢ MONG ĐỢI KHI ĐÃ CÓ CHỐT CHẶN:
    - Trạm 2 BẮT BUỘC PHẢI BÁO FAIL.
    - GATE_01 và GATE_05 BẮT BUỘC PHẢI LÀ FALSE.
    - error_traces phải ghi rõ: 'Bỏ sót ... thực thể ảnh gốc chưa được bóc tách: ...GT_SCHEME_1'.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        # 1. Tạo file ảnh duy nhất cho Fig 1
        img1 = _create_dummy_image_with_sidecar(tmpdir, "fig1.png", b"CONTENT_FIG_1")

        # 2. Ground Truth Manifest chứa 3 thực thể
        gt_manifest = {
            "source_document_path": "references/sample_paper.pdf",
            "section_id": "MUC_3_1",
            "total_entities": 3,
            "entities": [
                {
                    "entity_id": "GT_SCHEME_1",
                    "entity_type": "SCHEME",
                    "page_number": 2,
                    "bbox": [50.0, 100.0, 500.0, 400.0],
                    "caption_text": "Scheme 1. Schematic illustration of synthesis",
                    "label_extracted": "Scheme 1",
                },
                {
                    "entity_id": "GT_FIG_1_A",
                    "entity_type": "SUB_PANEL",
                    "parent_id": "GT_FIG_1",
                    "page_number": 3,
                    "bbox": [50.0, 100.0, 270.0, 350.0],
                    "caption_text": "Fig. 1. (a) XRD patterns",
                    "label_extracted": "Fig. 1(a)",
                },
                {
                    "entity_id": "GT_FIG_1_B",
                    "entity_type": "SUB_PANEL",
                    "parent_id": "GT_FIG_1",
                    "page_number": 3,
                    "bbox": [280.0, 100.0, 500.0, 350.0],
                    "caption_text": "Fig. 1. (b) FT-IR spectra",
                    "label_extracted": "Fig. 1(b)",
                },
            ],
            "scan_timestamp": "2026-10-25T00:00:00",
            "scanner_version": "1.0.0-pymupdf-dom",
        }

        # 3. Asset Manifest chỉ bóc tách 1 ảnh (bỏ sót GT_SCHEME_1 và GT_FIG_1_B)
        asset_item_data = {
            "asset_id": "FIG_3_1_1",
            "mpt_reference_id": "MPT_3_1_1",
            "asset_type": "DSAFE_ORIGINAL_CROP",
            "file_path_local": img1,
            "aspect_ratio": "16:9",
            "matlab_palette_colors_used": ["#0072BD"],
            "geometry_audit": {
                "envelope_clearance_percent": 10.0,
                "min_font_size_pt": 10.5,
                "layout_collision_detected": False,
                "dpi": 300,
                "file_size_bytes": 1024,
            },
            "pedagogical_caption": {
                "chain_1_origin": "XRD patterns",
                "chain_2_physical_feature": "Đỉnh nhiễu xạ Bragg",
                "chain_3_mathematical_link": "Phương trình Bragg",
                "chain_4_threshold_inflection": "Góc 2theta",
                "chain_5_engineering_insight": "Độ tinh khiết pha",
            },
            "ground_truth_ref": "GT_FIG_1_A",
        }
        manifest = VisualAssetManifest(
            section_id="MUC_3_1",
            total_assets=1,
            assets=[VisualAssetItem(**asset_item_data)],
        )

        # 4. Blueprint chỉ bao phủ GT_FIG_1_A
        blueprint = CurriculumBlueprint(
            course_name="Quang học và Quang điện tử",
            chapter_index=3,
            section_index=1,
            section_title="Quang xúc tác",
            source_documents=["references/sample_paper.pdf"],
            dag_nodes=[
                KnowledgeDAGNode(
                    node_id="NODE_1",
                    concept_title="Nhiễu xạ XRD",
                    bloom_level="Hiểu",
                    first_principles_root="Bragg Law",
                )
            ],
            dag_edges=[],
            mpt_specifications=[
                MicroPointTopicSpec(
                    mpt_id="MPT_3_1_1",
                    title="Phân tích cấu trúc tinh thể qua XRD",
                    core_proposition="Pha tinh thể quyết định hoạt tính quang xúc tác",
                    prerequisite_mpt_ids=[],
                    target_equations=["2d\\sin\\theta = n\\lambda"],
                    ground_truth_ref=["GT_FIG_1_A"],
                )
            ],
            terminology_contracts=[
                TerminologyEntry(
                    term_en="photocatalysis",
                    term_vi="quang xúc tác",
                    symbol_latex="\\eta",
                    standard_unit_si="\\%",
                    iso_standard_ref="ISO 80000",
                    decree_30_case="lower",
                )
            ],
        )

        draft_payload = _create_minimal_payload("MUC_3_1")

        # 5. Khởi tạo QualityGatekeeper và nạp đầu vào
        gatekeeper = QualityGatekeeper()
        gatekeeper.load_inputs(
            blueprint_input=blueprint,
            manifest_input=manifest,
            draft_payload_input=draft_payload,
            source_manifest_input=gt_manifest,
        )

        # 6. Chạy Trạm 2
        st2_result = gatekeeper.run_dohg_station_2_manifest_audit()

        # Kiểm chứng chốt chặn: Trạm 2 BẮT BUỘC PHẢI BÁO FAIL
        assert st2_result.passed is False, "Trạm 2 phải báo FAIL khi bỏ sót thực thể Ground Truth!"
        assert st2_result.error_traces is not None
        traces_text = " ".join(st2_result.error_traces)
        assert "GT_SCHEME_1" in traces_text, "Lỗi phải chỉ rõ bỏ sót GT_SCHEME_1!"

        # Kiểm chứng Bảng kiểm 17 điểm: GATE_01 và GATE_05 phải là False
        dummy_pass_st = StationResult(station_name="Dummy", passed=True, details="OK")
        dohg_res = DOHGStationsResult(
            station_1_static_linter=dummy_pass_st,
            station_2_boundary_re3_audit=st2_result,
            station_3_sympy_cas_verification=dummy_pass_st,
            station_4_typographic_preflight=dummy_pass_st,
            station_5_reportlab_compilation=dummy_pass_st,
        )
        b17 = gatekeeper.evaluate_binary_17_checklist(
            dohg_res,
            pub_meta=PublicationMetadata(
                temp_build_path="",
                final_drive_path="",
                atomic_replace_success=True,
                file_size_bytes=100,
                sha256_checksum="abc",
                two_pass_total_pages="Trang 1 / 1",
                pdf_bookmarks_levels_verified=True,
            ),
        )
        assert b17.checklist_status["GATE_01_elastic_depth_zero_omission"] is False
        assert b17.checklist_status["GATE_05_five_link_ocr_caption"] is False


def test_02_detect_sha256_duplicate_collision():
    """
    [BÀI TEST CHỐNG TRÙNG LẶP HÌNH ẢNH]:
    Mô phỏng việc nhúng cùng 1 file ảnh (hoặc ảnh có SHA-256 giống hệt)
    cho cả 2 slot asset khác nhau.
    KẾT QUẢ MONG ĐỢI:
    - Trạm 2 BẮT BUỘC PHẢI BÁO FAIL với thông báo: 'Trùng lặp mã băm SHA-256...'.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        duplicate_content = b"IDENTICAL_IMAGE_BINARY_DATA_FOR_COLLISION_TEST"
        img1 = _create_dummy_image_with_sidecar(tmpdir, "img_slot_1.png", duplicate_content)
        img2 = _create_dummy_image_with_sidecar(tmpdir, "img_slot_2.png", duplicate_content)

        manifest = VisualAssetManifest(
            section_id="MUC_3_1",
            total_assets=2,
            assets=[
                VisualAssetItem(
                    asset_id="ASSET_1",
                    mpt_reference_id="MPT_3_1_1",
                    asset_type="DSAFE_ORIGINAL_CROP",
                    file_path_local=img1,
                    aspect_ratio="16:9",
                    matlab_palette_colors_used=["#0072BD"],
                    geometry_audit=TripleBufferMetrics(
                        envelope_clearance_percent=10.0,
                        min_font_size_pt=10.5,
                        layout_collision_detected=False,
                        dpi=300,
                        file_size_bytes=1024,
                    ),
                    pedagogical_caption=OCRChainPedagogicalCaption(
                        chain_1_origin="Origin 1",
                        chain_2_physical_feature="Feature 1",
                        chain_3_mathematical_link="Link 1",
                        chain_4_threshold_inflection="Threshold 1",
                        chain_5_engineering_insight="Insight 1",
                    ),
                ),
                VisualAssetItem(
                    asset_id="ASSET_2",
                    mpt_reference_id="MPT_3_1_2",
                    asset_type="DSAFE_ORIGINAL_CROP",
                    file_path_local=img2,
                    aspect_ratio="16:9",
                    matlab_palette_colors_used=["#0072BD"],
                    geometry_audit=TripleBufferMetrics(
                        envelope_clearance_percent=10.0,
                        min_font_size_pt=10.5,
                        layout_collision_detected=False,
                        dpi=300,
                        file_size_bytes=1024,
                    ),
                    pedagogical_caption=OCRChainPedagogicalCaption(
                        chain_1_origin="Origin 2",
                        chain_2_physical_feature="Feature 2",
                        chain_3_mathematical_link="Link 2",
                        chain_4_threshold_inflection="Threshold 2",
                        chain_5_engineering_insight="Insight 2",
                    ),
                ),
            ],
        )

        blueprint = CurriculumBlueprint(
            course_name="Quang học và Quang điện tử",
            chapter_index=3,
            section_index=1,
            section_title="Quang xúc tác",
            source_documents=["references/sample_paper.pdf"],
            dag_nodes=[
                KnowledgeDAGNode(
                    node_id="NODE_1",
                    concept_title="Nhiễu xạ XRD",
                    bloom_level="Hiểu",
                    first_principles_root="Bragg Law",
                )
            ],
            dag_edges=[],
            mpt_specifications=[
                MicroPointTopicSpec(
                    mpt_id="MPT_3_1_1",
                    title="Luận điểm 1",
                    core_proposition="Mệnh đề 1",
                    prerequisite_mpt_ids=[],
                    target_equations=["E = mc^2"],
                )
            ],
            terminology_contracts=[
                TerminologyEntry(
                    term_en="photocatalysis",
                    term_vi="quang xúc tác",
                    symbol_latex="\\eta",
                    standard_unit_si="\\%",
                    iso_standard_ref="ISO 80000",
                    decree_30_case="lower",
                )
            ],
        )

        gatekeeper = QualityGatekeeper()
        gatekeeper.load_inputs(
            blueprint_input=blueprint,
            manifest_input=manifest,
            draft_payload_input=_create_minimal_payload("MUC_3_1"),
        )

        st2 = gatekeeper.run_dohg_station_2_manifest_audit()
        assert st2.passed is False, "Trạm 2 phải báo FAIL khi phát hiện 2 file ảnh trùng lặp SHA-256!"
        assert st2.error_traces is not None
        traces_text = " ".join(st2.error_traces)
        assert "Trùng lặp mã băm SHA-256" in traces_text or "trùng lặp" in traces_text.lower()


def test_03_zero_omission_full_coverage_pass():
    """
    [BÀI TEST NGHIỆM THU - CHUYỂN SANG PASS 100%]:
    Cung cấp đầy đủ 100% thực thể:
    - Bóc tách độc lập: Scheme 1, Fig 1a, Fig 1b.
    - Ánh xạ đủ MPT: MPT_3.1.0 (Scheme 1), MPT_3.1.1 (Fig 1a), MPT_3.1.2 (Fig 1b).
    KẾT QUẢ MONG ĐỢI:
    - Trạm 2 BÁO PASS 100%.
    - GATE_01 = True, GATE_05 = True.
    - Bảng kiểm 17 điểm đạt 17/17 PASS.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        img_scheme = _create_dummy_image_with_sidecar(tmpdir, "scheme1.png", b"UNIQUE_CONTENT_SCHEME_1")
        img_fig1a = _create_dummy_image_with_sidecar(tmpdir, "fig1a.png", b"UNIQUE_CONTENT_FIG_1A")
        img_fig1b = _create_dummy_image_with_sidecar(tmpdir, "fig1b.png", b"UNIQUE_CONTENT_FIG_1B")

        gt_manifest = {
            "source_document_path": "references/sample_paper.pdf",
            "section_id": "MUC_3_1",
            "total_entities": 3,
            "entities": [
                {
                    "entity_id": "GT_SCHEME_1",
                    "entity_type": "SCHEME",
                    "page_number": 2,
                    "bbox": [50.0, 100.0, 500.0, 400.0],
                    "caption_text": "Scheme 1. Synthesis",
                    "label_extracted": "Scheme 1",
                },
                {
                    "entity_id": "GT_FIG_1_A",
                    "entity_type": "SUB_PANEL",
                    "page_number": 3,
                    "bbox": [50.0, 100.0, 270.0, 350.0],
                    "caption_text": "Fig. 1. (a) XRD",
                    "label_extracted": "Fig. 1(a)",
                },
                {
                    "entity_id": "GT_FIG_1_B",
                    "entity_type": "SUB_PANEL",
                    "page_number": 3,
                    "bbox": [280.0, 100.0, 500.0, 350.0],
                    "caption_text": "Fig. 1. (b) FTIR",
                    "label_extracted": "Fig. 1(b)",
                },
            ],
            "scan_timestamp": "2026-10-25T00:00:00",
            "scanner_version": "1.0.0-pymupdf-dom",
        }

        manifest = VisualAssetManifest(
            section_id="MUC_3_1",
            total_assets=3,
            assets=[
                VisualAssetItem(
                    asset_id="ASSET_SCHEME",
                    mpt_reference_id="MPT_3_1_0",
                    asset_type="SCHEMATIC_VECTOR_DIAGRAM",
                    file_path_local=img_scheme,
                    aspect_ratio="16:9",
                    matlab_palette_colors_used=["#0072BD"],
                    geometry_audit=TripleBufferMetrics(
                        envelope_clearance_percent=12.0,
                        min_font_size_pt=10.5,
                        layout_collision_detected=False,
                        dpi=300,
                        file_size_bytes=1024,
                    ),
                    pedagogical_caption=OCRChainPedagogicalCaption(
                        chain_1_origin="Origin Scheme",
                        chain_2_physical_feature="Feature Scheme",
                        chain_3_mathematical_link="Link Scheme",
                        chain_4_threshold_inflection="Threshold Scheme",
                        chain_5_engineering_insight="Insight Scheme",
                    ),
                    ground_truth_ref="GT_SCHEME_1",
                ),
                VisualAssetItem(
                    asset_id="ASSET_FIG1A",
                    mpt_reference_id="MPT_3_1_1",
                    asset_type="DSAFE_ORIGINAL_CROP",
                    file_path_local=img_fig1a,
                    aspect_ratio="16:9",
                    matlab_palette_colors_used=["#0072BD"],
                    geometry_audit=TripleBufferMetrics(
                        envelope_clearance_percent=11.0,
                        min_font_size_pt=10.5,
                        layout_collision_detected=False,
                        dpi=300,
                        file_size_bytes=1024,
                    ),
                    pedagogical_caption=OCRChainPedagogicalCaption(
                        chain_1_origin="Origin 1A",
                        chain_2_physical_feature="Feature 1A",
                        chain_3_mathematical_link="Link 1A",
                        chain_4_threshold_inflection="Threshold 1A",
                        chain_5_engineering_insight="Insight 1A",
                    ),
                    ground_truth_ref="GT_FIG_1_A",
                ),
                VisualAssetItem(
                    asset_id="ASSET_FIG1B",
                    mpt_reference_id="MPT_3_1_2",
                    asset_type="DSAFE_ORIGINAL_CROP",
                    file_path_local=img_fig1b,
                    aspect_ratio="16:9",
                    matlab_palette_colors_used=["#0072BD"],
                    geometry_audit=TripleBufferMetrics(
                        envelope_clearance_percent=10.0,
                        min_font_size_pt=10.5,
                        layout_collision_detected=False,
                        dpi=300,
                        file_size_bytes=1024,
                    ),
                    pedagogical_caption=OCRChainPedagogicalCaption(
                        chain_1_origin="Origin 1B",
                        chain_2_physical_feature="Feature 1B",
                        chain_3_mathematical_link="Link 1B",
                        chain_4_threshold_inflection="Threshold 1B",
                        chain_5_engineering_insight="Insight 1B",
                    ),
                    ground_truth_ref="GT_FIG_1_B",
                ),
            ],
        )

        blueprint = CurriculumBlueprint(
            course_name="Quang học và Quang điện tử",
            chapter_index=3,
            section_index=1,
            section_title="Quang xúc tác",
            source_documents=["references/sample_paper.pdf"],
            dag_nodes=[
                KnowledgeDAGNode(
                    node_id="NODE_1",
                    concept_title="Nhiễu xạ XRD",
                    bloom_level="Hiểu",
                    first_principles_root="Bragg Law",
                )
            ],
            dag_edges=[],
            mpt_specifications=[
                MicroPointTopicSpec(
                    mpt_id="MPT_3_1_0",
                    title="Quy trình tổng hợp",
                    core_proposition="Tổng hợp thủy nhiệt",
                    prerequisite_mpt_ids=[],
                    target_equations=["T = 180^{\\circ}\\text{C}"],
                    ground_truth_ref=["GT_SCHEME_1"],
                ),
                MicroPointTopicSpec(
                    mpt_id="MPT_3_1_1",
                    title="Pha tinh thể XRD",
                    core_proposition="Đỉnh Bragg đặc trưng",
                    prerequisite_mpt_ids=["MPT_3_1_0"],
                    target_equations=["2d\\sin\\theta = n\\lambda"],
                    ground_truth_ref=["GT_FIG_1_A"],
                ),
                MicroPointTopicSpec(
                    mpt_id="MPT_3_1_2",
                    title="Liên kết hóa học FT-IR",
                    core_proposition="Dao động liên kết V-O và W-O",
                    prerequisite_mpt_ids=["MPT_3_1_1"],
                    target_equations=["\\nu = \\frac{1}{2\\pi c} \\sqrt{\\frac{k}{\\mu}}"],
                    ground_truth_ref=["GT_FIG_1_B"],
                ),
            ],
            terminology_contracts=[
                TerminologyEntry(
                    term_en="photocatalysis",
                    term_vi="quang xúc tác",
                    symbol_latex="\\eta",
                    standard_unit_si="\\%",
                    iso_standard_ref="ISO 80000",
                    decree_30_case="lower",
                )
            ],
        )

        gatekeeper = QualityGatekeeper()
        gatekeeper.load_inputs(
            blueprint_input=blueprint,
            manifest_input=manifest,
            draft_payload_input=_create_minimal_payload("MUC_3_1"),
            source_manifest_input=gt_manifest,
        )

        st2 = gatekeeper.run_dohg_station_2_manifest_audit()
        assert st2.passed is True, f"Trạm 2 phải PASS 100% khi đủ hết ảnh và MPT! Lỗi: {st2.error_traces}"

        dummy_pass_st = StationResult(station_name="Dummy", passed=True, details="OK")
        dohg_res = DOHGStationsResult(
            station_1_static_linter=dummy_pass_st,
            station_2_boundary_re3_audit=st2,
            station_3_sympy_cas_verification=dummy_pass_st,
            station_4_typographic_preflight=dummy_pass_st,
            station_5_reportlab_compilation=dummy_pass_st,
        )
        b17 = gatekeeper.evaluate_binary_17_checklist(
            dohg_res,
            pub_meta=PublicationMetadata(
                temp_build_path="",
                final_drive_path="",
                atomic_replace_success=True,
                file_size_bytes=100,
                sha256_checksum="abc",
                two_pass_total_pages="Trang 1 / 1",
                pdf_bookmarks_levels_verified=True,
            ),
        )
        assert b17.checklist_status["GATE_01_elastic_depth_zero_omission"] is True
        assert b17.checklist_status["GATE_05_five_link_ocr_caption"] is True
        assert b17.total_passed == 17


def test_04_docx_cross_audit_matrix_red_alert():
    """
    [BÀI TEST GIAO DIỆN DOCX]:
    Biên dịch hồ sơ kiểm toán .docx khi có thực thể bị bỏ sót:
    - Kiểm tra bảng Ma trận 3 cột xuất hiện trong tài liệu Word.
    - Kiểm tra có đoạn text 'BÁO ĐỘNG ĐỎ' và cell chứa 'OMITTED'.
    """
    import docx
    from scripts.docx_audit_generator import (
        CrossAuditEntry,
        SectionAuditItem,
        ChapterAuditPayload,
        DocxAuditGenerator,
    )

    with tempfile.TemporaryDirectory() as tmpdir:
        out_docx = os.path.join(tmpdir, "test_audit.docx")

        matrix = [
            CrossAuditEntry(
                ground_truth_entity="Scheme 1. Quy trình tổng hợp nanocomposite (Trang 2)",
                extracted_asset="[BỎ SÓT / OMITTED]",
                assigned_mpt="[CẢNH BÁO: THIẾU MPT]",
                status="OMITTED",
            ),
            CrossAuditEntry(
                ground_truth_entity="Fig. 1(a). Giản đồ XRD (Trang 3)",
                extracted_asset="[Goc]_Muc_3.1_Hinh_01a_XRD.png",
                assigned_mpt="MPT_3.1.1: Phân tích cấu trúc tinh thể XRD",
                status="PASS",
            ),
        ]

        sec = SectionAuditItem(
            section_id="3.1",
            section_title="Quang xúc tác xử lý môi trường",
            pdf_filename="Chuyen_Khao_Muc_3.1.pdf",
            cross_audit_matrix=matrix,
        )

        payload = ChapterAuditPayload(
            subject_name="Quang học và Quang điện tử",
            chapter_id="3",
            chapter_title="Vật liệu Quang xúc tác Bán dẫn",
            sections=[sec],
        )

        doc = DocxAuditGenerator.build_document(payload)
        doc.save(out_docx)

        assert os.path.exists(out_docx), "File .docx phải được sinh ra."

        doc = docx.Document(out_docx)

        # 1. Kiểm tra tồn tại Bảng Ma Trận Đối Chiếu Chéo 3 Cột
        table_found = False
        for tbl in doc.tables:
            header_row_text = " ".join([c.text for c in tbl.rows[0].cells])
            if "Thực Thể Nguồn Ground Truth" in header_row_text:
                table_found = True
                break
        assert table_found, "Bảng Ma trận đối chiếu chéo 3 cột phải hiện diện trong .docx!"

        # 2. Kiểm tra Banner Cảnh Báo Đỏ (Red Alert) khi tỷ lệ < 100%
        full_text = "\n".join([p.text for p in doc.paragraphs])
        for tbl in doc.tables:
            for row in tbl.rows:
                for c in row.cells:
                    full_text += "\n" + c.text

        assert "BÁO ĐỘNG ĐỎ" in full_text, "Phải xuất hiện cảnh báo 'BÁO ĐỘNG ĐỎ' khi tỷ lệ bao phủ < 100%!"
        assert "OMITTED" in full_text or "BỎ SÓT" in full_text


def test_05_ground_truth_manifest_engine_pdf_scan():
    """
    [BÀI TEST ĐỘNG CƠ QUÉT TẤT ĐỊNH PYMUPDF DOM]:
    Tạo một file PDF tổng hợp chứa Scheme 1, Fig 1 (panel a, b) và Table 1.
    Chạy GroundTruthManifestEngine.scan_document():
    - Xác minh phát hiện đầy đủ GT_SCHEME_1, GT_FIG_1, GT_FIG_1_A, GT_FIG_1_B, GT_TABLE_1.
    - Xác minh xuất bản và nạp lại JSON (Roundtrip).
    """
    import fitz
    from scripts.ground_truth_manifest_engine import (
        GroundTruthManifestEngine,
        GroundTruthManifest,
    )

    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, "synthetic_paper.pdf")
        json_path = os.path.join(tmpdir, "ground_truth.json")

        # 1. Tạo PDF giả lập bằng PyMuPDF
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)

        page.insert_text(
            fitz.Point(50, 150),
            "Scheme 1. Schematic illustration of synthesis route for ternary nanocomposite AVWGC.",
            fontsize=11,
        )
        page.insert_text(
            fitz.Point(50, 350),
            "Fig. 1. Characterization of samples: (a) XRD patterns and (b) FT-IR spectra.",
            fontsize=11,
        )
        page.insert_text(
            fitz.Point(50, 550),
            "Table 1. Elemental composition of synthesized nanocomposites.",
            fontsize=11,
        )
        doc.save(pdf_path)
        doc.close()

        # 2. Quét DOM qua GroundTruthManifestEngine
        engine = GroundTruthManifestEngine(pdf_path)
        manifest = engine.scan_document(section_id="MUC_3_1")

        assert isinstance(manifest, GroundTruthManifest)
        assert manifest.total_entities >= 4

        entity_ids = {e.entity_id for e in manifest.entities}
        assert "GT_SCHEME_1" in entity_ids, "Phải nhận diện được GT_SCHEME_1"
        assert "GT_FIG_1_A" in entity_ids, "Phải nhận diện được sub-panel GT_FIG_1_A"
        assert "GT_FIG_1_B" in entity_ids, "Phải nhận diện được sub-panel GT_FIG_1_B"
        assert "GT_TABLE_1" in entity_ids, "Phải nhận diện được GT_TABLE_1"

        # 3. Xuất JSON và kiểm tra roundtrip
        manifest.export_json(json_path)
        assert os.path.exists(json_path)

        restored = GroundTruthManifest.load_json(json_path)
        assert restored.total_entities == manifest.total_entities
        assert len(restored.entities) == len(manifest.entities)

